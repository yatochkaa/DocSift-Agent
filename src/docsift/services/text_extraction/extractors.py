from __future__ import annotations

import csv
import logging
from pathlib import Path

import fitz
import openpyxl
import xlrd
from PIL import Image

from docsift.schemas.text_extraction import (
    BoundingBox,
    ExtractedPage,
    ExtractedTable,
    TextBlock,
)
from docsift.services.text_extraction.ocr import OcrEngineProtocol
from docsift.services.text_extraction.preprocessing import ImagePreprocessor

logger = logging.getLogger(__name__)


def has_usable_text_layer(text: str, word_count: int) -> bool:
    visible = [character for character in text if not character.isspace()]
    if not visible:
        return False
    meaningful_count = sum(character.isalnum() for character in visible)
    readable_count = sum(
        character.isprintable() and character != "\ufffd" for character in visible
    )
    readable_ratio = readable_count / len(visible)
    return meaningful_count >= 50 and word_count >= 10 and readable_ratio >= 0.8


def _bbox(
    x0: float,
    y0: float,
    x1: float,
    y1: float,
    width: float,
    height: float,
) -> BoundingBox:
    return BoundingBox(
        x0=max(0.0, min(1.0, x0 / width)),
        y0=max(0.0, min(1.0, y0 / height)),
        x1=max(0.0, min(1.0, x1 / width)),
        y1=max(0.0, min(1.0, y1 / height)),
    )


class PdfTextExtractor:
    def __init__(
        self,
        ocr_engine: OcrEngineProtocol,
        preprocessor: ImagePreprocessor,
        render_dpi: int = 300,
    ) -> None:
        self._ocr_engine = ocr_engine
        self._preprocessor = preprocessor
        self._render_dpi = render_dpi

    def extract(self, path: Path) -> list[ExtractedPage]:
        pages: list[ExtractedPage] = []
        with fitz.open(path) as document:
            if document.page_count == 0:
                raise ValueError("PDF contains no pages")
            for page_index, page in enumerate(document):
                text = page.get_text("text")
                words = page.get_text("words")
                if has_usable_text_layer(text, len(words)):
                    pages.append(self._extract_text_page(page, page_index + 1))
                else:
                    pages.append(self._extract_ocr_page(page, page_index + 1))
        return pages

    def _extract_text_page(self, page: fitz.Page, number: int) -> ExtractedPage:
        width = float(page.rect.width)
        height = float(page.rect.height)
        blocks: list[TextBlock] = []
        for raw_block in page.get_text("blocks"):
            if len(raw_block) >= 7 and raw_block[6] != 0:
                continue
            text = str(raw_block[4]).strip()
            if not text:
                continue
            blocks.append(
                TextBlock(
                    text=text,
                    bbox=_bbox(
                        float(raw_block[0]),
                        float(raw_block[1]),
                        float(raw_block[2]),
                        float(raw_block[3]),
                        width,
                        height,
                    ),
                    confidence=1.0,
                    source="pdf:text_layer",
                )
            )

        return ExtractedPage(
            number=number,
            width=width,
            height=height,
            blocks=blocks,
            tables=self._extract_tables(page, width, height),
            used_ocr=False,
        )

    def _extract_ocr_page(self, page: fitz.Page, number: int) -> ExtractedPage:
        scale = self._render_dpi / 72
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB)
        rendered = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        prepared = self._preprocessor.prepare(rendered)
        blocks, tables = self._ocr_engine.extract(prepared)
        return ExtractedPage(
            number=number,
            width=float(prepared.width),
            height=float(prepared.height),
            blocks=blocks,
            tables=tables,
            used_ocr=True,
        )

    @staticmethod
    def _extract_tables(page: fitz.Page, width: float, height: float) -> list[ExtractedTable]:
        try:
            found_tables = page.find_tables().tables
        except Exception:
            logger.exception("PyMuPDF table detection failed on page %s", page.number + 1)
            return []

        tables: list[ExtractedTable] = []
        for table in found_tables:
            rows = [
                ["" if cell is None else str(cell).strip() for cell in row]
                for row in table.extract()
            ]
            rows = [row for row in rows if any(row)]
            if not rows:
                continue
            x0, y0, x1, y1 = table.bbox
            tables.append(
                ExtractedTable(
                    rows=rows,
                    bbox=_bbox(x0, y0, x1, y1, width, height),
                    source="pdf:table",
                )
            )
        return tables


class ImageTextExtractor:
    def __init__(
        self,
        ocr_engine: OcrEngineProtocol,
        preprocessor: ImagePreprocessor,
    ) -> None:
        self._ocr_engine = ocr_engine
        self._preprocessor = preprocessor

    def extract(self, path: Path) -> list[ExtractedPage]:
        with Image.open(path) as source:
            prepared = self._preprocessor.prepare(source)
        blocks, tables = self._ocr_engine.extract(prepared)
        return [
            ExtractedPage(
                number=1,
                width=float(prepared.width),
                height=float(prepared.height),
                blocks=blocks,
                tables=tables,
                used_ocr=True,
            )
        ]


class SpreadsheetTextExtractor:
    def extract(self, path: Path) -> list[ExtractedPage]:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            sheets = [(path.stem, self._read_csv(path))]
        elif suffix == ".xlsx":
            sheets = self._read_xlsx(path)
        elif suffix == ".xls":
            sheets = self._read_xls(path)
        else:
            raise ValueError(f"Unsupported spreadsheet type: {suffix}")

        return [
            self._page_from_rows(number, label, rows)
            for number, (label, rows) in enumerate(sheets, start=1)
        ]

    @staticmethod
    def _read_csv(path: Path) -> list[list[str]]:
        content = path.read_bytes()
        text: str | None = None
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = content.decode("utf-8", errors="replace")

        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[cell.strip() for cell in row] for row in csv.reader(text.splitlines(), dialect)]

    @staticmethod
    def _read_xlsx(path: Path) -> list[tuple[str, list[list[str]]]]:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                (
                    worksheet.title,
                    [
                        ["" if value is None else str(value) for value in row]
                        for row in worksheet.iter_rows(values_only=True)
                    ],
                )
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()

    @staticmethod
    def _read_xls(path: Path) -> list[tuple[str, list[list[str]]]]:
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            return [
                (
                    sheet.name,
                    [
                        ["" if value is None else str(value) for value in sheet.row_values(row)]
                        for row in range(sheet.nrows)
                    ],
                )
                for sheet in workbook.sheets()
            ]
        finally:
            workbook.release_resources()

    @staticmethod
    def _page_from_rows(number: int, label: str, rows: list[list[str]]) -> ExtractedPage:
        row_count = max(1, len(rows))
        column_count = max(1, max((len(row) for row in rows), default=0))
        blocks: list[TextBlock] = []

        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                text = value.strip()
                if not text:
                    continue
                blocks.append(
                    TextBlock(
                        text=text,
                        bbox=BoundingBox(
                            x0=column_index / column_count,
                            y0=row_index / row_count,
                            x1=(column_index + 1) / column_count,
                            y1=(row_index + 1) / row_count,
                        ),
                        confidence=1.0,
                        source=f"spreadsheet:{label}",
                    )
                )

        tables = (
            [
                ExtractedTable(
                    rows=rows,
                    bbox=BoundingBox(x0=0.0, y0=0.0, x1=1.0, y1=1.0),
                    source=f"spreadsheet:{label}",
                )
            ]
            if rows
            else []
        )
        return ExtractedPage(
            number=number,
            label=label,
            width=float(column_count),
            height=float(row_count),
            blocks=blocks,
            tables=tables,
            used_ocr=False,
        )

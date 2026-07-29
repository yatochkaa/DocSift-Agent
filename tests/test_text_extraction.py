from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw

from docsift.schemas.text_extraction import BoundingBox, ExtractedTable, TextBlock
from docsift.services.text_extraction import TextExtractionService
from docsift.services.text_extraction.extractors import PdfTextExtractor
from docsift.services.text_extraction.preprocessing import ImagePreprocessor


class StubOcrEngine:
    def extract(self, image: Image.Image) -> tuple[list[TextBlock], list[ExtractedTable]]:
        return (
            [
                TextBlock(
                    text="Распознанный текст документа",
                    bbox=BoundingBox(x0=0.1, y0=0.1, x1=0.9, y1=0.2),
                    confidence=0.95,
                    source="ocr:stub",
                )
            ],
            [],
        )


def _make_photo(path: Path, *, rotate: bool = False) -> Path:
    image = Image.new("RGB", (800, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 100), "INVOICE 12345", fill="black")
    draw.text((80, 180), "TOTAL 1000 VAT 200", fill="black")
    if rotate:
        image = image.rotate(2, expand=True, fillcolor="white")
    image.save(path)
    return path


def test_text_pdf_uses_embedded_text_layer(tmp_path: Path) -> None:
    path = tmp_path / "text-document.pdf"
    document = fitz.open()
    page = document.new_page()
    text = (
        "Invoice number 123 Supplier Alpha Buyer Beta Total amount 1000 VAT 200 "
        "Date 2026 07 24 payment goods services contract delivery address"
    )
    page.insert_textbox(fitz.Rect(50, 50, 550, 300), text, fontsize=12)
    document.save(path)
    document.close()

    result = TextExtractionService(ocr_engine=StubOcrEngine()).extract(path)

    assert result.used_ocr is False
    assert result.coordinate_space == "normalized"
    assert result.pages[0].blocks
    assert result.pages[0].blocks[0].source == "pdf:text_layer"


def test_scanned_pdf_renders_page_and_uses_ocr(tmp_path: Path) -> None:
    image_path = _make_photo(tmp_path / "scan-source.png")
    image_bytes = BytesIO()
    with Image.open(image_path) as image:
        image.save(image_bytes, format="PNG")

    path = tmp_path / "scanned-document.pdf"
    document = fitz.open()
    page = document.new_page(width=800, height=500)
    page.insert_image(page.rect, stream=image_bytes.getvalue())
    document.save(path)
    document.close()

    result = TextExtractionService(ocr_engine=StubOcrEngine()).extract(path)

    assert result.used_ocr is True
    assert result.pages[0].used_ocr is True
    assert result.pages[0].blocks
    assert result.pages[0].blocks[0].source == "ocr:stub"


def test_photo_is_preprocessed_and_uses_ocr(tmp_path: Path) -> None:
    path = _make_photo(tmp_path / "chat-photo.jpg", rotate=True)

    result = TextExtractionService(ocr_engine=StubOcrEngine()).extract(path)

    assert result.used_ocr is True
    assert result.media_type == "image/jpeg"
    assert result.pages[0].blocks
    assert result.pages[0].width > 0
    assert result.pages[0].height > 0


def test_csv_is_read_directly_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "invoice.csv"
    path.write_text("name;quantity;amount\nТовар;2;1000\n", encoding="utf-8")

    result = TextExtractionService(ocr_engine=StubOcrEngine()).extract(path)

    assert result.used_ocr is False
    assert result.pages[0].blocks
    assert result.pages[0].tables[0].rows[1] == ["Товар", "2", "1000"]


def test_xlsx_is_read_directly_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "invoice.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Счёт"
    worksheet.append(["Наименование", "Количество", "Сумма"])
    worksheet.append(["Услуга", 1, 2500])
    workbook.save(path)
    workbook.close()

    result = TextExtractionService(ocr_engine=StubOcrEngine()).extract(path)

    assert result.used_ocr is False
    assert result.pages[0].label == "Счёт"
    assert result.pages[0].blocks
    assert result.pages[0].tables[0].rows[1] == ["Услуга", "1", "2500"]


# ── PDF render-limit tests ──────────────────────────────────────────────


class StubOcrEngineEmpty:
    """OCR stub returning no blocks — enough to satisfy the protocol."""

    def extract(self, image: Image.Image) -> tuple[list[TextBlock], list[ExtractedTable]]:
        return [], []


class RecordingPreprocessor:
    """Preprocessor stub that records the image size and passes it through."""

    def __init__(self) -> None:
        self.last_size: tuple[int, int] | None = None

    def prepare(self, image: Image.Image) -> Image.Image:
        self.last_size = image.size
        return image.convert("RGB")


def test_oversized_page_scale_is_reduced(tmp_path: Path) -> None:
    """A 2000×2000 pt page at 300 DPI would be 8333×8333 px (~69 MP).
    With a 9 MP limit the scale must be reduced (new_scale ≈ 1.5),
    producing ~3000×3000 px — within budget, no exception."""
    max_mp = 9
    render_dpi = 300
    page_size = 2000

    path = tmp_path / "oversized.pdf"
    doc = fitz.open()
    doc.new_page(width=page_size, height=page_size)
    doc.save(path)
    doc.close()

    preprocessor = RecordingPreprocessor()
    extractor = PdfTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=preprocessor,  # type: ignore[arg-type]
        render_dpi=render_dpi,
        max_render_megapixels=max_mp,
    )

    pages = extractor.extract(path)

    assert len(pages) == 1
    assert preprocessor.last_size is not None
    w, h = preprocessor.last_size
    megapixels = w * h / 1_000_000
    # Allow 1% tolerance for integer rounding in PyMuPDF pixmap dimensions
    assert megapixels <= max_mp * 1.01, f"Rendered {megapixels:.1f} MP exceeds {max_mp} MP limit"
    assert w > 0 and h > 0

    # Verify scale was actually reduced, not just clamped to 1.0
    observed_scale = w / page_size
    assert 1.0 < observed_scale < render_dpi / 72, (
        f"Expected reduced scale, got {observed_scale:.4f}"
    )


def test_huge_mediabox_page_rejected(tmp_path: Path) -> None:
    """A 14400×14400 pt page at default 40 MP limit — even at 72 DPI (207 MP)
    exceeds the budget. ValueError must be raised before any pixmap allocation."""
    path = tmp_path / "huge_mediabox.pdf"
    doc = fitz.open()
    doc.new_page(width=14400, height=14400)
    doc.save(path)
    doc.close()

    extractor = PdfTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=RecordingPreprocessor(),  # type: ignore[arg-type]
    )

    with pytest.raises(ValueError, match="превышает лимит 40 МП"):
        extractor.extract(path)


def test_normal_a4_renders_without_scale_reduction(tmp_path: Path) -> None:
    """A normal A4 page (595×842 pt) at 300 DPI should render at ~2479×3508 px
    with no scale reduction."""
    render_dpi = 300
    expected_w = round(595 * render_dpi / 72)  # 2479
    expected_h = round(842 * render_dpi / 72)  # 3508

    path = tmp_path / "a4.pdf"
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    doc.save(path)
    doc.close()

    preprocessor = RecordingPreprocessor()
    extractor = PdfTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=preprocessor,  # type: ignore[arg-type]
        render_dpi=render_dpi,
    )

    pages = extractor.extract(path)

    assert len(pages) == 1
    assert preprocessor.last_size is not None
    w, h = preprocessor.last_size
    assert abs(w - expected_w) <= 2, f"Width {w} not within ±2 of {expected_w}"
    assert abs(h - expected_h) <= 2, f"Height {h} not within ±2 of {expected_h}"


def test_too_many_pages_raises_value_error(tmp_path: Path) -> None:
    """A PDF with more pages than max_pages raises ValueError."""
    path = tmp_path / "three_pages.pdf"
    doc = fitz.open()
    for _ in range(3):
        doc.new_page()
    doc.save(path)
    doc.close()

    extractor = PdfTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=RecordingPreprocessor(),  # type: ignore[arg-type]
        max_pages=2,
    )

    with pytest.raises(ValueError, match="допустимый лимит"):
        extractor.extract(path)


def test_service_passes_pdf_limits_to_extractor() -> None:
    """TextExtractionService forwards pdf_max_pages / pdf_max_render_megapixels."""
    service = TextExtractionService(pdf_max_pages=7, pdf_max_render_megapixels=11)

    assert service._pdf._max_pages == 7  # type: ignore[attr-defined]
    assert service._pdf._max_render_megapixels == 11  # type: ignore[attr-defined]


def test_service_defaults_match_settings_defaults() -> None:
    """Service defaults equal Settings defaults — catches drift between the two."""
    from docsift.core.config import Settings

    defaults = Settings()
    service = TextExtractionService()

    assert service._pdf._max_pages == defaults.pdf_max_pages  # type: ignore[attr-defined]
    assert service._pdf._max_render_megapixels == defaults.pdf_max_render_megapixels  # type: ignore[attr-defined]
    assert (
        service._image._max_megapixels
        == Settings.model_fields["image_max_megapixels"].default
    )


# ── Image pixel-limit tests ────────────────────────────────────────────


def test_image_within_pixel_limit_is_processed(tmp_path: Path) -> None:
    """A 800×600 (0.48 MP) image with max_megapixels=1 is accepted."""
    from docsift.services.text_extraction.extractors import ImageTextExtractor

    image = Image.new("RGB", (800, 600), "white")
    image.save(tmp_path / "small.png")

    extractor = ImageTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=RecordingPreprocessor(),  # type: ignore[arg-type]
        max_megapixels=1,
    )

    pages = extractor.extract(tmp_path / "small.png")

    assert len(pages) == 1
    assert pages[0].used_ocr is True


def test_image_exceeding_pixel_limit_is_rejected(tmp_path: Path) -> None:
    """A 1200×1200 (1.44 MP) image with max_megapixels=1 raises before processing."""
    from docsift.services.text_extraction.extractors import ImageTextExtractor

    image = Image.new("RGB", (1200, 1200), "white")
    image.save(tmp_path / "big.png")

    preprocessor = RecordingPreprocessor()
    extractor = ImageTextExtractor(
        ocr_engine=StubOcrEngineEmpty(),
        preprocessor=preprocessor,  # type: ignore[arg-type]
        max_megapixels=1,
    )

    with pytest.raises(ValueError, match="превышает лимит 1 МП"):
        extractor.extract(tmp_path / "big.png")

    assert preprocessor.last_size is None


def test_service_passes_image_limit_to_extractor() -> None:
    """TextExtractionService forwards image_max_megapixels to ImageTextExtractor."""
    service = TextExtractionService(image_max_megapixels=13)

    assert service._image._max_megapixels == 13  # type: ignore[attr-defined]
